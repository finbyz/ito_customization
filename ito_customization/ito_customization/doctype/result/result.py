import re

import frappe
from frappe.model.document import Document
from frappe import _
from frappe.utils import cint, flt


class Result(Document):
	def autoname(self):
		if not self.student_name:
			frappe.throw("Student Name is required")

		self.name = self.student_name

	def validate(self):
		self.sync_round_1_values()

	def sync_round_1_values(self):
		"""Keep Round 1 score and topper in the matching Round 2 row."""
		round_2_by_subject = {
			row.subject_2: row
			for row in self.round_2 or []
			if row.subject_2
		}

		for round_1_row in self.round_1 or []:
			if not round_1_row.subject:
				continue

			round_2_row = round_2_by_subject.get(round_1_row.subject)
			if not round_2_row and self.qualified:
				round_2_row = self.append("round_2", {"subject_2": round_1_row.subject})
				round_2_by_subject[round_1_row.subject] = round_2_row

			if round_2_row:
				round_2_row.round_1_score = round_1_row.score
				round_2_row.round_1_topper = round_1_row.topper or ""


@frappe.whitelist()
def upload_results(file_url):
    """
    Import Result documents from uploaded CSV/Excel file.

    Rows are grouped by roll_no so multiple subjects belonging
    to the same student are stored in the same Result document.
    """

    if not file_url:
        frappe.throw(_("Please upload a result file."))

    file_doc = frappe.get_doc("File", {"file_url": file_url})
    file_path = file_doc.get_full_path()

    rows = read_result_file(file_path)

    if not rows:
        frappe.throw(_("The uploaded file does not contain any data."))

    grouped_results = group_results(rows)

    result = {
        "created": 0,
        "updated": 0,
        "rows_processed": len(rows),
        "errors": [],
    }

    for roll_no, student_rows in grouped_results.items():
        try:
            import_student_result(
                roll_no,
                student_rows,
                result,
            )

        except Exception:
            frappe.log_error(
                frappe.get_traceback(),
                f"Result Import Error - {roll_no}",
            )

            result["errors"].append(
                f"{roll_no}: {frappe.get_traceback().splitlines()[-1]}"
            )

    return result


# ---------------------------------------------------------------------
# FILE READING
# ---------------------------------------------------------------------

def read_result_file(file_path):
    import os
    import pandas as pd

    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".csv":
        dataframe = pd.read_csv(file_path)

    elif extension in [".xlsx", ".xls"]:
        dataframe = pd.read_excel(file_path)

    else:
        frappe.throw(
            _("Only CSV and Excel files are supported.")
        )

    dataframe = dataframe.fillna("")

    return dataframe.to_dict("records")


# ---------------------------------------------------------------------
# GROUP RESULTS
# ---------------------------------------------------------------------

def group_results(rows):
    """
    Group all rows belonging to the same student.

    Multiple subjects for the same roll number will therefore
    be stored in one Result document.
    """

    grouped = {}

    for row in rows:
        roll_no = clean_value(
            row.get("roll_no")
        )

        if not roll_no:
            continue

        grouped.setdefault(
            roll_no,
            []
        ).append(row)

    return grouped


# ---------------------------------------------------------------------
# IMPORT STUDENT RESULT
# ---------------------------------------------------------------------

def import_student_result(roll_no, rows, result):
    first_row = rows[0]

    existing_name = frappe.db.exists(
        "Result",
        {
            "roll_no": roll_no,
        },
    )

    if existing_name:
        doc = frappe.get_doc(
            "Result",
            existing_name
        )

        # Rebuild child tables during re-import
        doc.round_1 = []
        doc.round_2 = []

        result["updated"] += 1

    else:
        doc = frappe.new_doc("Result")

        result["created"] += 1

    # -------------------------------------------------------------
    # Parent fields
    # -------------------------------------------------------------

        # Keeping school code exactly as received for now.
    doc.school = clean_value(
        first_row.get("school_code")
    )

    doc.student_name = clean_value(
        first_row.get("student_name")
    )

    doc.parent_name = clean_value(
        first_row.get("parent_name")
    )

    doc.roll_no = roll_no

    doc.qualified = cint(
        first_row.get("qualified_round_two") or 0
    )

    doc.set(
    	"class",
    	clean_value(first_row.get("std_id"))
    )

    doc.mobile = clean_value(
        first_row.get("mobile")
    )

    doc.email = clean_value(
        first_row.get("email_id")
    )
    # -------------------------------------------------------------
    # Child tables
    # -------------------------------------------------------------

    for row in rows:
        add_round_1_row(
            doc,
            row
        )

        add_round_2_row(
            doc,
            row
        )

    doc.save(
        ignore_permissions=True
    )


# ---------------------------------------------------------------------
# ROUND 1
# ---------------------------------------------------------------------

def add_round_1_row(doc, row):
    subject = clean_value(
        row.get("subject_id")
    )

    academic_year = clean_value(
        row.get("academic_year")
    )

    score = row.get("round_1_score")
    rank = row.get("round_1_rank")

    # Don't create an empty Round 1 row.
    if not any([
        subject,
        academic_year,
        score != "",
        rank != "",
        row.get("round_1_award"),
        row.get("round_1_topper"),
        row.get("round_1_message"),
    ]):
        return

    child = doc.append(
        "round_1",
        {}
    )

    # -------------------------------------------------------------
    # Dynamic Subject Mapping
    # -------------------------------------------------------------

    child.subject = get_subject_from_excel(
        subject
    )

    # Academic year intentionally kept as-is.
    child.academic_year = academic_year

    if score != "":
        child.score = flt(score)

    if rank != "":
        child.rank = cint(rank)

    child.award = clean_value(
        row.get("round_1_award")
    )

    child.topper = clean_value(
        row.get("round_1_topper")
    )

    child.message = clean_value(
        row.get("round_1_message")
    )


# ---------------------------------------------------------------------
# ROUND 2
# ---------------------------------------------------------------------

def add_round_2_row(doc, row):
    subject = clean_value(
        row.get("subject_id")
    )

    academic_year = clean_value(
        row.get("academic_year")
    )

    score = row.get("round_2_score")
    rank = row.get("round_2_rank")

    # Don't create an empty Round 2 row.
    if not any([
        subject,
        score != "",
        rank != "",
        row.get("round_2_award"),
        row.get("round_2_topper"),
        row.get("round_2_message"),
    ]):
        return

    child = doc.append(
        "round_2",
        {}
    )

    # -------------------------------------------------------------
    # Dynamic Subject Mapping
    # -------------------------------------------------------------

    child.subject_2 = get_subject_from_excel(
        subject
    )

    # Academic year intentionally kept as-is.
    child.academic_year_2 = academic_year

    if score != "":
        child.score_2 = clean_value(
            row.get("round_2_score")
        )

    if rank != "":
        child.rank_2 = clean_value(
            row.get("round_2_rank")
        )

    child.award_2 = clean_value(
        row.get("round_2_award")
    )

    child.topper_2 = clean_value(
        row.get("round_2_topper")
    )

    child.message_2 = clean_value(
        row.get("round_2_message")
    )


# ---------------------------------------------------------------------
# DYNAMIC SUBJECT MAPPING
# ---------------------------------------------------------------------

def get_subject_from_excel(excel_subject):

    excel_subject = clean_value(
        excel_subject
    )

    if not excel_subject:
        return None

    excel_normalized = normalize_subject(
        excel_subject
    )

    subjects = frappe.get_all(
        "School Subject",
        fields=["name"],
    )

    matches = []

    for subject in subjects:
        subject_name = clean_value(
            subject.get("name")
        )

        subject_normalized = normalize_subject(
            subject_name
        )

        # Exact normalized match
        if excel_normalized == subject_normalized:
            matches.append(subject_name)
            continue

        # Excel subject contained inside School Subject name
        if excel_normalized in subject_normalized:
            matches.append(subject_name)
            continue

        # School Subject name contained inside Excel subject
        if subject_normalized in excel_normalized:
            matches.append(subject_name)

    # Remove duplicates while maintaining order.
    matches = list(dict.fromkeys(matches))

    if len(matches) == 1:
        return matches[0]

    if len(matches) > 1:
        frappe.throw(
            _(
                "Multiple School Subjects matched '{0}': {1}"
            ).format(
                excel_subject,
                ", ".join(matches),
            )
        )

    frappe.throw(
        _(
            "Could not map Excel subject '{0}' "
            "to any School Subject."
        ).format(
            excel_subject
        )
    )


# ---------------------------------------------------------------------
# SUBJECT NORMALIZATION
# ---------------------------------------------------------------------

def normalize_subject(value):

    value = clean_value(value).lower()

    # Remove olympiad-related wording.
    value = re.sub(
        r"\bolympiad\b",
        "",
        value,
    )

    # Remove brackets and their contents.
    value = re.sub(
        r"\([^)]*\)",
        "",
        value,
    )

    # Normalize common subject variations.
    replacements = {
        "mathematics": "maths",
        "math": "maths",
    }

    for old_value, new_value in replacements.items():
        value = re.sub(
            rf"\b{re.escape(old_value)}\b",
            new_value,
            value,
        )

    # Replace non-alphanumeric characters with spaces.
    value = re.sub(
        r"[^a-z0-9]+",
        " ",
        value,
    )

    # Remove extra spaces.
    value = " ".join(
        value.split()
    )

    return value


# ---------------------------------------------------------------------
# HELPERS
# ---------------------------------------------------------------------

def clean_value(value):
    if value is None:
        return ""

    return str(value).strip()


@frappe.whitelist()
def download_result_template():
    """
    Generate and download the Result import template.
    """

    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    headers = [
        "school_code",
        "student_name",
        "parent_name",
        "roll_no",
        "std_id",
        "mobile",
        "email_id",
        "qualified_round_two",
        "subject_id",
        "academic_year",
        "round_1_score",
        "round_1_rank",
        "round_1_award",
        "round_1_topper",
        "round_1_message",
        "round_2_score",
        "round_2_rank",
        "round_2_award",
        "round_2_topper",
        "round_2_message",
    ]

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Result Import"

    # ---------------------------------------------------------
    # Headers
    # ---------------------------------------------------------

    for column_number, header in enumerate(headers, start=1):
        cell = worksheet.cell(
            row=1,
            column=column_number,
            value=header,
        )

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    # ---------------------------------------------------------
    # Example / Instructions Sheet
    # ---------------------------------------------------------

    instructions = workbook.create_sheet(
        "Instructions"
    )

    instruction_rows = [
		["Result Import Template"],
		[""],
		["Column", "Description"],

		[
			"school_code",
			"School code",
		],
		[
			"student_name",
			"Student name",
		],
		[
			"parent_name",
			"Parent name",
		],
		[
			"roll_no",
			"Student roll number",
		],
		[
			"std_id",
			"Student class",
		],
		[
			"mobile",
			"Student/parent mobile number",
		],
		[
			"email_id",
			"Student/parent email",
		],
		[
			"qualified_round_two",
			"1 for qualified, 0 for not qualified",
		],
		[
			"subject_id",
			"Subject name as used in the source data",
		],
		[
			"academic_year",
			'Please enter Academic Year in this format: "AY-2025/26"',
		],
		[
			"round_1_score",
			"Round 1 score",
		],
		[
			"round_1_rank",
			"Round 1 rank",
		],
		[
			"round_1_award",
			"Round 1 award",
		],
		[
			"round_1_topper",
			"Round 1 topper",
		],
		[
			"round_1_message",
			"Round 1 message",
		],
		[
			"round_2_score",
			"Round 2 score",
		],
		[
			"round_2_rank",
			"Round 2 rank",
		],
		[
			"round_2_award",
			"Round 2 award",
		],
		[
			"round_2_topper",
			"Round 2 topper",
		],
		[
			"round_2_message",
			"Round 2 message",
		],
	]

    for row_number, row in enumerate(
        instruction_rows,
        start=1,
    ):
        for column_number, value in enumerate(
            row,
            start=1,
        ):
            instructions.cell(
                row=row_number,
                column=column_number,
                value=value,
            )

    # ---------------------------------------------------------
    # Formatting
    # ---------------------------------------------------------

    instructions["A1"].font = Font(
        bold=True,
        size=14,
    )

    instructions["A3"].font = Font(
        bold=True
    )

    instructions["B3"].font = Font(
        bold=True
    )

    worksheet.freeze_panes = "A2"

    for column in worksheet.columns:
        max_length = 0

        column_letter = column[0].column_letter

        for cell in column:
            if cell.value:
                max_length = max(
                    max_length,
                    len(str(cell.value)),
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 3,
            40,
        )

    instructions.column_dimensions[
        "A"
    ].width = 30

    instructions.column_dimensions[
        "B"
    ].width = 60

    # ---------------------------------------------------------
    # Return Excel File
    # ---------------------------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    frappe.response["filename"] = (
        "Result_Import_Template.xlsx"
    )

    frappe.response["filecontent"] = (
        output.getvalue()
    )

    frappe.response["type"] = "download"
